"""Supplier summary and title detail for the filtered operation export."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .payables_priority import AGES, decorate_suppliers


HEADERS = (
    "Prioridade", "Fornecedor", "Código do fornecedor", "Criticidade",
    "Total em atraso do fornecedor", "Maior atraso do fornecedor (dias)",
    "Pagar agora do fornecedor", "Saldo a negociar do fornecedor",
    "Título", "Parcela", "Documento", "Emissão", "Lançamento", "Vencimento",
    "Dias em atraso", "Valor", "Saldo a pagar", "Status", "Descrição",
    "Nível de criticidade", "Classificação", "Faixa de atraso do título",
)


SUMMARY_HEADERS = (
    "Nome Fornecedor", "Valor Total Vencido", "Vencimento da última NF",
    "Valor da última NF", "Criticidade", "Caducidade", "Quadrante",
)


def _money(value):
    return Decimal(str(value or 0))


def _date(value):
    return date.fromisoformat(str(value)[:10]) if value else None


def _latest_invoice(supplier):
    """Latest due date among open titles in the applied filters."""
    titles = [
        title for title in supplier.get("titulos", [])
        if title.get("data_vencimento") and _money(title.get("saldo_a_pagar")) > 0
    ]
    return max(titles, key=lambda title: (
        _date(title["data_vencimento"]),
        _date(title.get("data_emissao")) or date.min,
        int(title.get("codigo_contas_pagar") or 0),
        int(title.get("codigo_parcela") or 0),
    ), default=None)


def _supplier_summary(supplier):
    title = _latest_invoice(supplier)
    level = supplier.get("nivel_criticidade")
    classification = (
        f"{level} - {supplier['criticidade_label']}"
        if level else supplier["criticidade_label"]
    )
    age_label, quadrant = "Não informado", "Não informado"
    if title:
        days = int(title.get("dias_vencidos") or 0)
        age = next((n for threshold, n in (
            (120, 5), (90, 4), (60, 3), (30, 2), (0, 1)
        ) if days > threshold), 0)
        age_label = f"{age} - {dict(AGES)[age]}" if age else "Sem atraso"
        quadrant = (
            f"Q{(5 - level) * 5 + (6 - age)}" if level and age
            else "Fora da matriz" if not age else "Classificação pendente"
        )
    invoice_value = title.get("valor_nota_fiscal") if title else None
    return [
        supplier.get("nome_fornecedor") or "",
        _money(supplier.get("valor_total_vencido")),
        _date(title["data_vencimento"]) if title else None,
        _money(invoice_value) if invoice_value is not None else None,
        classification, age_label, quadrant,
    ]


def _add_supplier_summary(workbook, suppliers):
    sheet = workbook.create_sheet("Resumo por fornecedor", 0)
    sheet.append(SUMMARY_HEADERS)
    for supplier in suppliers:
        sheet.append(_supplier_summary(supplier))
    for row in sheet.iter_rows(min_row=2, max_row=len(suppliers) + 1):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[1].number_format = row[3].number_format = '"R$" #,##0.00'
        row[2].number_format = "dd/mm/yyyy"
        sheet.row_dimensions[row[0].row].height = 16 * max(
            2, len(wrap(str(row[0].value or ""), width=44))
        )
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B7088")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 42
    for column, width in enumerate((48, 24, 25, 24, 24, 24, 22), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    workbook.active = 0


def build_payables_workbook(fornecedores):
    fornecedores = decorate_suppliers(fornecedores)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Contas vencidas"
    sheet.append(HEADERS)
    row_number = 1
    for priority, supplier in enumerate(fornecedores, 1):
        for title in supplier.get("titulos", []):
            balance = _money(title.get("saldo_a_pagar"))
            days = int(title.get("dias_vencidos") or 0)
            sheet.append([
                priority, supplier.get("nome_fornecedor") or "",
                str(supplier.get("codigo_fornecedor") or ""),
                "Crítico" if supplier.get("critico") else "Não crítico",
                _money(supplier.get("valor_total_vencido")),
                int(supplier.get("dias_atraso") or 0),
                _money(supplier.get("pagamento_imediato")),
                _money(supplier.get("saldo_negociar")),
                str(title.get("codigo_parcela") or ""),
                str(title.get("numero_parcela") or ""),
                str(title.get("numero_documento") or ""),
                _date(title.get("data_emissao")),
                _date(title.get("data_lancamento")),
                _date(title.get("data_vencimento")), days,
                _money(title.get("valor_total")), balance,
                "Pago" if balance <= 0 else "Em atraso" if days > 0 else "A vencer",
                title.get("descricao_conta") or "",
                supplier.get("nivel_criticidade"), supplier["criticidade_label"],
                next((label for level, label in AGES if days > {5:120, 4:90, 3:60, 2:30, 1:0}[level]), "Sem atraso"),
            ])
            row_number += 1
            for column in range(1, len(HEADERS) + 1):
                cell = sheet.cell(row=row_number, column=column)
                # API text must stay literal, including values beginning with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.column in (5, 7, 8, 16, 17):
                    cell.number_format = '"R$" #,##0.00'
                elif cell.column in (12, 13, 14):
                    cell.number_format = "dd/mm/yyyy"
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B7088")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 42
    for column in range(1, len(HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = (
            60 if column == 19 else 36 if column == 2 else 22
        )
    _add_supplier_summary(workbook, fornecedores)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
