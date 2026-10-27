from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook

from core.payables_export import build_payables_workbook
from core.payables_priority import decorate_suppliers
from core.tests import ContasPagarTests
from core.test_payables_priority import PriorityTests


class UnregisteredSupplierTests(TestCase):
    setUp = PriorityTests.setUp

    def payload(self):
        payload = ContasPagarTests.payload()
        supplier = payload['fornecedores'][0]
        supplier.update(codigo_fornecedor='sem-cadastro-bm9tZTpDU1JG',
                        nome_fornecedor='CSRF', sem_cadastro=True,
                        critico=False, nivel_criticidade=None)
        return payload

    def test_no_automatic_classification_or_internal_id_in_excel(self):
        supplier = self.payload()['fornecedores'][0]
        decorated = decorate_suppliers([supplier])[0]
        self.assertIsNone(decorated['nivel_criticidade'])
        workbook = load_workbook(BytesIO(build_payables_workbook([supplier])))
        self.assertEqual(workbook.active['E2'].value, 'Sem classifica\u00e7\u00e3o')
        self.assertEqual(workbook.active['G2'].value, 'Classifica\u00e7\u00e3o pendente')
        sheet = workbook['Contas vencidas']
        self.assertEqual(sheet['C2'].value, 'Sem cadastro no MV')
        self.assertEqual(sheet['D2'].value, 'Sem classifica\u00e7\u00e3o')
        self.assertIsNone(sheet['T2'].value)

    @patch('core.views.api_get')
    def test_unregistered_titles_visible_without_payment_forms(self, api_get):
        api_get.return_value = self.payload()
        response = self.client.get('/financeiro/contas-a-pagar/operacao/')
        self.assertContains(response, 'CSRF')
        self.assertContains(response, 'Sem cadastro no MV')
        self.assertContains(response, 'Medicamentos')
        self.assertNotContains(response, 'name="form_action" value="pagamento_salvar"')
        self.assertNotContains(response, 'name="form_action" value="pagamento_atualizar"')
        self.assertNotContains(response, 'name="form_action" value="pagamento_excluir"')

    @patch('core.views.api_get')
    def test_unregistered_supplier_has_classification_form(self, api_get):
        api_get.return_value = self.payload()
        response = self.client.get('/administrativo/fornecedores-criticos/')
        self.assertContains(response, 'CSRF')
        self.assertNotContains(response, 'Cadastre o fornecedor no MV')
        self.assertContains(response, 'Salvar classifica\u00e7\u00e3o')

    @patch('core.views.api_post')
    @patch('core.views.api_patch')
    def test_synthetic_classification_uses_separate_route_but_payments_remain_blocked(self, api_patch, api_post):
        key = self.payload()['fornecedores'][0]['codigo_fornecedor']
        self.client.post('/administrativo/fornecedores-criticos/', {
            'codigo_fornecedor': key, 'nivel_criticidade': '4', 'justificativa': 'Teste',
        })
        self.client.post('/financeiro/contas-a-pagar/operacao/', {
            'codigo_fornecedor': key, 'form_action': 'pagamento_salvar',
            'codigo_parcela': '99', 'valor_pago': '100', 'data_pagamento': '2026-10-27',
            'banco': 'Teste', 'agencia': '1', 'numero_conta': '1',
        })
        api_patch.assert_called_once_with(
            '/app_glosas/financeiro/contas-a-pagar/sem-cadastro/' + key + '/criticidade',
            {'nivel_criticidade': 4, 'justificativa': 'Teste'},
        )
        api_post.assert_not_called()

    def test_explicit_level_survives_decoration_and_excel(self):
        supplier = self.payload()['fornecedores'][0]
        supplier.update(nivel_criticidade=5, critico=True, classificado_por='Financeiro')
        book = load_workbook(BytesIO(build_payables_workbook([supplier])))
        self.assertEqual(supplier['nivel_criticidade'], 5)
        self.assertEqual(book.active['E2'].value, '5 - Extrema')
        self.assertEqual(book.active['G2'].value, 'Q2')
        self.assertEqual(book['Contas vencidas']['D2'].value, 'Crítico')
        self.assertEqual(book['Contas vencidas']['C2'].value, 'Sem cadastro no MV')

    @patch('core.views.api_patch')
    def test_unregistered_classification_requires_level_and_valid_identity(self, api_patch):
        for fields in [
            {'codigo_fornecedor': 'sem-cadastro-bm9tZTpDU1JG', 'critico': '1'},
            {'codigo_fornecedor': 'sem-cadastro-../../x', 'nivel_criticidade': '5'},
        ]:
            self.client.post('/administrativo/fornecedores-criticos/', fields)
        api_patch.assert_not_called()
