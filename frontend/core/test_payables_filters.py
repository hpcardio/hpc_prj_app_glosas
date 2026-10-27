from copy import deepcopy
from io import BytesIO
from urllib.parse import parse_qs, urlparse
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook

from core.tests import ContasPagarTests
from core.test_payables_priority import PriorityTests
from core.payables_priority import decorate_matrix


class PayablesFilterTests(TestCase):
    setUp = PriorityTests.setUp

    def payload(self):
        payload = deepcopy(ContasPagarTests.payload())
        payload['opcoes_fornecedores'] = [
            {'codigo': '10', 'nome': 'Fornecedor A'},
            {'codigo': '20', 'nome': 'Fornecedor B'},
        ]
        return payload

    @patch('core.views.api_get')
    def test_selected_names_chips_and_filter_state_survive_pagination(self, api_get):
        payload = self.payload()
        payload.update(total=40, total_pages=2)
        api_get.return_value = payload
        response = self.client.get('/financeiro/contas-a-pagar/operacao/', {
            'fornecedores_ids': '10,20', 'ordenar': 'maior_valor_vencido', 'faixa_atraso': '1',
        })
        self.assertContains(response, 'Fornecedor A')
        self.assertContains(response, 'Fornecedor B')
        self.assertContains(response, 'data-supplier-chips')
        self.assertNotContains(response, '<select name="fornecedores_ids"')
        self.assertContains(response, 'Caducidade')
        self.assertNotContains(response, 'Atraso do título')
        self.assertContains(response, 'Maior valor vencido')
        params = api_get.call_args.kwargs['params']
        self.assertEqual(params['fornecedores_ids'], '10,20')
        self.assertEqual(params['incluir_opcoes'], 'true')
        for link in [response.context['pagination']['next_url'], response.context['export_url']]:
            query = parse_qs(urlparse(link).query)
            self.assertEqual(query['fornecedores_ids'], ['10,20'])
            self.assertEqual(query['ordenar'], ['maior_valor_vencido'])

    @patch('core.views.api_get')
    def test_management_replaces_binary_criticality_with_quadrant(self, api_get):
        api_get.side_effect = [self.payload(), {'linhas': [], 'colunas': [], 'pendentes': {'quantidade': 0}}]
        response = self.client.get('/financeiro/contas-a-pagar/gestao/', {
            'fornecedores_ids': '10,20', 'quadrante': '5', 'criticidade': 'nao_criticos',
            'ordenar': 'maior_valor_vencido',
        })
        self.assertContains(response, 'name="quadrante"')
        self.assertContains(response, '>Q25</option>')
        self.assertNotContains(response, '<select class="form-select" name="criticidade">')
        self.assertContains(response, 'name="nivel_criticidade"')
        for call in api_get.call_args_list:
            self.assertEqual(call.kwargs['params']['fornecedores_ids'], '10,20')
            self.assertEqual(call.kwargs['params']['quadrante'], '5')
            self.assertEqual(call.kwargs['params']['criticidade'], 'todos')

    @patch('core.views.api_get')
    def test_export_keeps_accumulated_ids_quadrant_and_order(self, api_get):
        api_get.return_value = self.payload()
        response = self.client.get('/financeiro/contas-a-pagar/operacao/', {
            'fornecedores_ids': '10,20', 'quadrante': '5',
            'ordenar': 'maior_valor_vencido', 'export': 'xlsx',
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(api_get.call_args.kwargs['params']['fornecedores_ids'], '10,20')
        self.assertEqual(api_get.call_args.kwargs['params']['quadrante'], '5')
        self.assertEqual(api_get.call_args.kwargs['params']['ordenar'], 'maior_valor_vencido')
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 2)

    def test_matrix_navigation_replaces_quadrant_but_keeps_suppliers(self):
        matrix = {'linhas': [{'celulas': [{'nivel': 4, 'faixa': 1}]}]}
        decorate_matrix(matrix, {'fornecedores_ids': '10,20', 'quadrante': '5', 'ordenar': 'maior_valor_vencido'})
        query = parse_qs(urlparse(matrix['linhas'][0]['celulas'][0]['url']).query)
        self.assertNotIn('quadrante', query)
        self.assertEqual(query['nivel_criticidade'], ['4'])
        self.assertEqual(query['fornecedores_ids'], ['10,20'])
