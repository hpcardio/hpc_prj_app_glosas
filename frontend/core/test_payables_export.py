from copy import deepcopy
from datetime import datetime
from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook

from .payables_export import HEADERS, SUMMARY_HEADERS, build_payables_workbook
from .services import ApiError
from . import tests as existing_tests


class PayablesExportTests(TestCase):
    setUp = existing_tests.ContasPagarTests.setUp
    url = "/financeiro/contas-a-pagar/operacao/"

    @patch("core.views.api_get")
    def test_exports_all_suppliers_in_one_request_with_applied_filters(self, api_get):
        first = existing_tests.ContasPagarTests.payload()
        first["total_pages"] = 1
        second = deepcopy(first)
        second["fornecedores"][0]["codigo_fornecedor"] = 20
        second["fornecedores"][0]["titulos"][0]["codigo_parcela"] = 100
        first["fornecedores"].extend(second["fornecedores"])
        api_get.return_value = first
        filters = dict(q="Essencial", criticidade="criticos", status="NEGOCIACAO",
                       data_inicio="2026-01-01", data_fim="2026-12-31")
        response = self.client.get(self.url, {**filters, "page": 9, "export": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(response["Cache-Control"], "private, no-store")
        workbook = load_workbook(BytesIO(response.content))
        summary = workbook.active
        self.assertEqual(tuple(c.value for c in summary[1]), SUMMARY_HEADERS)
        self.assertEqual(summary.max_row, 3)
        self.assertEqual(summary["B2"].value, 500000)
        sheet = workbook["Contas vencidas"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(tuple(c.value for c in sheet[1]), HEADERS)
        self.assertEqual(sheet["A3"].value, 2)
        self.assertEqual(sheet["C3"].value, "20")
        self.assertEqual(sheet["I3"].value, "100")
        self.assertEqual(sheet["L2"].value, datetime(2026, 5, 18))
        self.assertEqual(sheet["P2"].value, 650000)
        self.assertEqual(sheet["R2"].value, "Em atraso")
        self.assertEqual(sheet["S2"].value, "Medicamentos")
        self.assertEqual(sheet["P2"].data_type, "n")
        api_get.assert_called_once()
        self.assertEqual(api_get.call_args.kwargs["params"], {**filters, "exportar": "true"})

    def test_literal_text_statuses_and_missing_dates(self):
        supplier = existing_tests.ContasPagarTests.payload()["fornecedores"][0]
        supplier["nome_fornecedor"] = "=1+1"
        title = supplier["titulos"][0]
        title.update(numero_documento="000123", descricao_conta="=HYPERLINK(x)",
                     data_emissao=None, saldo_a_pagar=0)
        future = deepcopy(title)
        future.update(saldo_a_pagar=100, dias_vencidos=0)
        supplier["titulos"].append(future)
        workbook = load_workbook(BytesIO(build_payables_workbook([supplier])))
        self.assertEqual(workbook.active["A2"].data_type, "s")
        sheet = workbook["Contas vencidas"]
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["S2"].data_type, "s")
        self.assertEqual(sheet["K2"].value, "000123")
        self.assertIsNone(sheet["L2"].value)
        self.assertEqual(sheet["R2"].value, "Pago")
        self.assertEqual(sheet["R3"].value, "A vencer")

    @patch("core.views.api_get")
    def test_empty_export_has_headers(self, api_get):
        api_get.return_value = {"fornecedores": [], "total_pages": 1}
        response = self.client.get(self.url, {"export": "xlsx"})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 1)

    @patch("core.views.api_get")
    def test_api_failure_never_downloads_partial_workbook(self, api_get):
        payload = existing_tests.ContasPagarTests.payload()
        payload["total_pages"] = 2
        api_get.return_value = payload
        response = self.client.get(self.url, {"q": "Essencial", "export": "xlsx"})
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("export=", response.url)
        self.assertIn("q=Essencial", response.url)

    @patch("core.views.api_get")
    def test_link_preserves_applied_query_without_current_page(self, api_get):
        payload = existing_tests.ContasPagarTests.payload()
        payload["total_pages"] = 3
        api_get.return_value = payload
        response = self.client.get(self.url, {"q": "A & B", "page": 2})
        self.assertContains(response, "Exportar Excel")
        self.assertIn("q=A+%26+B", response.context["export_url"])
        self.assertNotIn("page=", response.context["export_url"])

    @patch("core.views.api_get")
    def test_export_requires_screen_permission(self, api_get):
        session = self.client.session
        session["api_user"] = {"perfil": "usuario", "telas_permitidas": []}
        session.save()
        response = self.client.get(self.url, {"export": "xlsx"})
        self.assertEqual(response.status_code, 302)
        api_get.assert_not_called()

    @patch("core.views.api_get")
    def test_export_requires_login(self, api_get):
        self.client.session.flush()
        response = self.client.get(self.url, {"export": "xlsx"})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith("/login/"))
        api_get.assert_not_called()

    @patch("core.views.api_get")
    def test_api_error_preserves_filters_without_download(self, api_get):
        api_get.side_effect = ApiError("Falha na API")
        response = self.client.get(self.url, {"q": "Essencial", "export": "xlsx"})
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("export=", response.url)
        self.assertIn("q=Essencial", response.url)


class SupplierSummaryExportTests(TestCase):
    def supplier(self):
        supplier = existing_tests.ContasPagarTests.payload()["fornecedores"][0]
        supplier["nivel_criticidade"] = 4
        supplier["titulos"][0]["valor_nota_fiscal"] = "1300000.00"
        return supplier

    def test_latest_due_date_drives_invoice_value_age_and_quadrant(self):
        supplier = self.supplier()
        older = supplier["titulos"][0]
        older["data_emissao"] = "2026-09-15"
        recent = {**older, "codigo_parcela": 100, "data_emissao": "2026-01-01",
                  "data_vencimento": "2026-09-17", "dias_vencidos": 30,
                  "valor_total": "100.00", "valor_nota_fiscal": "300.00"}
        paid = {**recent, "codigo_parcela": 101, "data_vencimento": "2026-12-01",
                "saldo_a_pagar": "0.00"}
        supplier["titulos"] += [recent, paid]
        workbook = load_workbook(BytesIO(build_payables_workbook([supplier])))
        summary = workbook.active
        self.assertEqual(summary.max_column, 7)
        self.assertEqual(summary.max_row, 2)
        self.assertEqual([c.value for c in summary[2]], [
            "Fornecedor Essencial", 500000, datetime(2026, 9, 17), 300,
            "4 - Alta", "1 - 1–30 dias", "Q10",
        ])
        self.assertEqual(summary["D2"].data_type, "n")
        self.assertEqual(summary["C2"].number_format, "dd/mm/yyyy")
        self.assertEqual(summary.auto_filter.ref, "A1:G2")
        self.assertEqual(workbook["Contas vencidas"].max_row, 4)

    def test_age_boundaries_and_pending_classification(self):
        cases = [(0, "Fora da matriz"), (1, "Q15"), (30, "Q15"),
                 (31, "Q14"), (60, "Q14"), (61, "Q13"), (90, "Q13"),
                 (91, "Q12"), (120, "Q12"), (121, "Q11")]
        for days, quadrant in cases:
            with self.subTest(days=days):
                supplier = self.supplier()
                supplier["nivel_criticidade"] = 3
                supplier["titulos"][0]["dias_vencidos"] = days
                sheet = load_workbook(BytesIO(build_payables_workbook([supplier]))).active
                self.assertEqual(sheet["G2"].value, quadrant)
        supplier["nivel_criticidade"] = None
        sheet = load_workbook(BytesIO(build_payables_workbook([supplier]))).active
        self.assertEqual(sheet["G2"].value, "Classificação pendente")

    def test_missing_invoice_data_is_not_invented(self):
        supplier = self.supplier()
        supplier["titulos"][0].pop("valor_nota_fiscal")
        sheet = load_workbook(BytesIO(build_payables_workbook([supplier]))).active
        self.assertIsNone(sheet["D2"].value)
        supplier["titulos"][0]["data_vencimento"] = None
        sheet = load_workbook(BytesIO(build_payables_workbook([supplier]))).active
        self.assertIsNone(sheet["C2"].value)
        self.assertEqual(sheet["G2"].value, "Não informado")
        supplier["titulos"] = []
        sheet = load_workbook(BytesIO(build_payables_workbook([supplier]))).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["B2"].value, 500000)

    def test_installments_do_not_multiply_original_invoice_value(self):
        supplier = self.supplier()
        title = supplier["titulos"][0]
        supplier["titulos"].append({**title, "codigo_parcela": 100,
                                   "data_vencimento": "2026-07-09"})
        sheet = load_workbook(BytesIO(build_payables_workbook([supplier]))).active
        self.assertEqual(sheet["D2"].value, 1300000)
        self.assertEqual(sheet["C2"].value, datetime(2026, 7, 9))
